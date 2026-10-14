import os
import csv
import openpyxl

CAMINHO_DIR = os.path.dirname(os.path.abspath(__file__))
CAMINHO_CSV = os.path.join(CAMINHO_DIR, "catalogo.csv")
CAMINHO_XLSX = os.path.join(CAMINHO_DIR, "catalogo.xlsx")


def carregar_produtos():
    """
    Carrega os produtos da planilha. Prioriza catalogo.csv para que você
    possa ler e editar o arquivo diretamente pelo VS Code sem caracteres estranhos!
    """
    if os.path.exists(CAMINHO_CSV):
        return _carregar_produtos_csv()
    elif os.path.exists(CAMINHO_XLSX):
        return _carregar_produtos_xlsx()
    else:
        criar_planilhas_padrao()
        return _carregar_produtos_csv()


def _carregar_produtos_csv():
    produtos_dict = {}
    try:
        with open(CAMINHO_CSV, mode="r", encoding="utf-8-sig") as f:
            primeira_linha = f.readline()
            delimitador = ";" if ";" in primeira_linha else ","
            f.seek(0)

            reader = csv.reader(f, delimiter=delimitador)
            header = next(reader, None)  # Pula o cabeçalho

            for row in reader:
                if not row or len(row) < 3 or not str(row[0]).strip():
                    continue

                nome_original = str(row[0]).strip()

                # Preço
                try:
                    p_raw = str(row[1]).replace("R$", "").replace(" ", "").replace(",", ".")
                    preco = float(p_raw)
                except (ValueError, TypeError):
                    preco = 0.0

                # Estoque
                try:
                    estoque = int(str(row[2]).strip())
                except (ValueError, TypeError):
                    estoque = 0

                chave = nome_original.lower()
                produtos_dict[chave] = {
                    "produto": nome_original,
                    "preco": preco,
                    "estoque": estoque
                }
    except Exception as erro:
        print(f"Erro ao ler catalogo.csv: {erro}")

    return produtos_dict


def _carregar_produtos_xlsx():
    produtos_dict = {}
    try:
        wb = openpyxl.load_workbook(CAMINHO_XLSX, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            nome_original = str(row[0]).strip()
            if not nome_original:
                continue
            try:
                preco_raw = row[1]
                if isinstance(preco_raw, str):
                    preco_raw = preco_raw.replace("R$", "").replace(" ", "").replace(",", ".")
                preco = float(preco_raw) if preco_raw is not None else 0.0
            except (ValueError, TypeError):
                preco = 0.0
            try:
                estoque = int(row[2]) if row[2] is not None else 0
            except (ValueError, TypeError):
                estoque = 0

            chave = nome_original.lower()
            produtos_dict[chave] = {
                "produto": nome_original,
                "preco": preco,
                "estoque": estoque
            }
        wb.close()
    except Exception as erro:
        print(f"Erro ao ler catalogo.xlsx: {erro}")
    return produtos_dict


def atualizar_estoque(nome, novo_estoque):
    """
    Atualiza o estoque do produto no catalogo.csv (e também no catalogo.xlsx se existir).
    """
    atualizado_csv = False

    if os.path.exists(CAMINHO_CSV):
        try:
            linhas = []
            delimitador = ";"
            with open(CAMINHO_CSV, mode="r", encoding="utf-8-sig") as f:
                primeira_linha = f.readline()
                delimitador = ";" if ";" in primeira_linha else ","
                f.seek(0)
                reader = csv.reader(f, delimiter=delimitador)
                for row in reader:
                    linhas.append(row)

            nome_busca = nome.strip().lower()
            for i in range(1, len(linhas)):
                if linhas[i] and len(linhas[i]) >= 3 and linhas[i][0].strip().lower() == nome_busca:
                    linhas[i][2] = str(int(novo_estoque))
                    atualizado_csv = True
                    break

            if atualizado_csv:
                with open(CAMINHO_CSV, mode="w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f, delimiter=delimitador)
                    writer.writerows(linhas)
                print(f"Estoque de '{nome}' atualizado para {novo_estoque} no catalogo.csv.")
        except PermissionError:
            print("AVISO: catalogo.csv está aberto em outro programa. Feche-o para salvar o estoque.")
        except Exception as erro:
            print(f"Erro ao atualizar catalogo.csv: {erro}")

    # Atualiza também o xlsx se existir
    if os.path.exists(CAMINHO_XLSX):
        try:
            wb = openpyxl.load_workbook(CAMINHO_XLSX)
            ws = wb.active
            nome_busca = nome.strip().lower()
            for row in range(2, ws.max_row + 1):
                val_produto = ws.cell(row=row, column=1).value
                if val_produto and str(val_produto).strip().lower() == nome_busca:
                    ws.cell(row=row, column=3, value=int(novo_estoque))
                    wb.save(CAMINHO_XLSX)
                    break
            wb.close()
        except Exception:
            pass

    return atualizado_csv


def criar_planilhas_padrao():
    """Cria o catalogo padrão."""
    itens = [
        ["Produto", "Preço", "Estoque"],
        ["Camisa Nossa Senhora", "45.00", "10"],
        ["Camisa São Bento", "50.00", "5"],
        ["Camisa Sagrado Coração", "48.00", "8"],
        ["Camisa São Miguel Arcanjo", "52.00", "12"]
    ]
    with open(CAMINHO_CSV, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerows(itens)


class CatalogoDinamico(dict):
    """
    Dicionário dinâmico que sempre reflete as atualizações da planilha
    sem necessidade de reiniciar o bot.
    """
    def __getitem__(self, item):
        return carregar_produtos()[item]

    def __iter__(self):
        return iter(carregar_produtos())

    def __len__(self):
        return len(carregar_produtos())

    def __contains__(self, item):
        return item in carregar_produtos()

    def get(self, item, default=None):
        return carregar_produtos().get(item, default)

    def items(self):
        return carregar_produtos().items()

    def values(self):
        return carregar_produtos().values()

    def keys(self):
        return carregar_produtos().keys()


# Mantém compatibilidade com `from catalogo import produtos`
produtos = CatalogoDinamico()